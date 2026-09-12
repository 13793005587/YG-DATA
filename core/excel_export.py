"""Excel 导出（通用）。

关键约束：
  * 表头取**所有行键的并集**（保持首次出现顺序），而不是只看第一行——
    否则首行缺少的字段会被静默丢弃。
  * 写入失败（文件被 Excel 占用、目录只读、磁盘满、路径过长）统一抛出
    :class:`ExportError`，由 UI 转成可操作的提示，绝不让程序崩溃。
"""

import logging

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

MAX_COLUMN_WIDTH = 40


class ExportError(RuntimeError):
    """导出失败，消息可直接展示给用户。"""


def _resolve_headers(data, headers):
    """确定表头：显式传入则用传入值，否则取全部行键的并集。"""
    if headers is not None:
        return list(headers)

    resolved = []
    seen = set()
    for row in data:
        for key in row:
            if key not in seen:
                seen.add(key)
                resolved.append(key)
    return resolved


def export_to_excel(data, file_path, headers=None, sheet_title="结果"):
    """
    将 ``list[dict]`` 导出为 xlsx 文件。

    data:    [{"列名1": 值1, "列名2": 值2, ...}, ...]
    headers: 可选，显式指定列与顺序（内部列可用它排除）
    """
    if not data:
        return

    headers = _resolve_headers(data, headers)
    if not headers:
        return

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title

    # 表头样式
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4A90E2")
    center = Alignment(horizontal="center", vertical="center")
    thin = Side(style="thin", color="D0D0D0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # 写表头
    for col_idx, name in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border

    # 写数据
    for row_idx, row in enumerate(data, start=2):
        for col_idx, key in enumerate(headers, start=1):
            value = row.get(key, "")
            if isinstance(value, str):
                value = value.strip()
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = center
            cell.border = border

    # 列宽自适应（按内容长度估算；中文按 2 个字符宽度计）
    for col_idx, key in enumerate(headers, start=1):
        max_len = _display_width(str(key))
        for row in data:
            v = str(row.get(key, "") or "")
            max_len = max(max_len, _display_width(v))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(
            max_len + 4, MAX_COLUMN_WIDTH
        )

    # 冻结表头
    ws.freeze_panes = "A2"

    try:
        wb.save(file_path)
    except PermissionError as e:
        raise ExportError(
            f"无法写入文件：\n{file_path}\n\n"
            "该文件可能正在 Excel 中打开，请关闭后重试；"
            "或换一个没有写入限制的位置。"
        ) from e
    except OSError as e:
        raise ExportError(f"导出失败：\n{file_path}\n\n{e}") from e
    except Exception as e:  # noqa: BLE001 - 交给 UI 统一提示，同时记录日志
        logger.exception("导出 Excel 失败：%s", file_path)
        raise ExportError(f"导出失败：\n{file_path}\n\n{e}") from e

    logger.info("已导出 %d 行到 %s", len(data), file_path)


def _display_width(text):
    """粗略计算显示宽度：CJK/全角字符算 2 列。"""
    width = 0
    for ch in text:
        width += 2 if ord(ch) > 0x2E7F else 1
    return width
