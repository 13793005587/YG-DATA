"""Excel 导出测试：列并集、失败防护、空数据。"""

import os
import shutil
import sys
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import load_workbook  # noqa: E402

from core.excel_export import ExportError, export_to_excel  # noqa: E402
from tests import make_temp_dir  # noqa: E402


class TestExportToExcel(unittest.TestCase):
    def setUp(self):
        self.tmpdir = make_temp_dir("export-")
        self.path = os.path.join(self.tmpdir, "out.xlsx")

    def tearDown(self):
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def _headers(self, path):
        wb = load_workbook(path)
        ws = wb.active
        headers = [c.value for c in ws[1]]
        wb.close()
        return headers

    def test_writes_headers_and_rows(self):
        data = [{"样品": "A", "表面能": 24.9}, {"样品": "B", "表面能": 30.1}]
        export_to_excel(data, self.path)
        self.assertEqual(self._headers(self.path), ["样品", "表面能"])

    def test_headers_are_union_of_all_rows(self):
        """回归：首行缺少的字段曾被静默丢弃。"""
        data = [{"样品": "A"}, {"样品": "B", "表面能": 1.0}]
        export_to_excel(data, self.path)
        self.assertEqual(self._headers(self.path), ["样品", "表面能"])

    def test_explicit_headers_control_columns_and_order(self):
        data = [{"序号": "A", "样品": "A", "表面能": 1.0}]
        export_to_excel(data, self.path, headers=["样品", "表面能"])
        self.assertEqual(self._headers(self.path), ["样品", "表面能"])

    def test_missing_values_become_empty_cells(self):
        data = [{"a": 1}, {"b": 2}]
        export_to_excel(data, self.path)
        wb = load_workbook(self.path)
        ws = wb.active
        self.assertEqual([c.value for c in ws[2]], [1, None])
        self.assertEqual([c.value for c in ws[3]], [None, 2])
        wb.close()

    def test_empty_data_writes_nothing(self):
        export_to_excel([], self.path)
        self.assertFalse(os.path.exists(self.path))

    def test_invalid_directory_raises_export_error(self):
        """回归：以前这里会抛出未捕获的 FileNotFoundError 导致程序崩溃。"""
        bad = os.path.join(self.tmpdir, "no", "such", "dir", "x.xlsx")
        with self.assertRaises(ExportError) as ctx:
            export_to_excel([{"a": 1}], bad)
        self.assertIn("导出失败", str(ctx.exception))

    def test_directory_path_raises_export_error(self):
        with self.assertRaises(ExportError):
            export_to_excel([{"a": 1}], self.tmpdir)
    def test_string_values_are_stripped(self):
        export_to_excel([{"a": "  x  "}], self.path)
        wb = load_workbook(self.path)
        self.assertEqual(wb.active.cell(row=2, column=1).value, "x")
        wb.close()

    def test_freeze_panes_and_column_width(self):
        export_to_excel([{"样品": "A"}], self.path)
        wb = load_workbook(self.path)
        ws = wb.active
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertGreater(ws.column_dimensions["A"].width, 2)
        wb.close()


if __name__ == "__main__":
    unittest.main()
